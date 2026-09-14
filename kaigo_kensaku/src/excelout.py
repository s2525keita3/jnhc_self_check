"""Excel出力（C案：サービス別の統合シート ＋ 市区町村別シート）。

ヘッダーは旧ツールと同じ3段構成（1行目:大見出し / 2行目:中見出し /
3行目:項目名、4行目からデータ）。先頭に「市区町村」列を追加しているので、
統合シートのままフィルタ・ピボットで分析できる。
"""
from __future__ import annotations

import logging
import os
import re
from typing import Dict, List

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

log = logging.getLogger(__name__)

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
GROUP_FILL = PatternFill("solid", fgColor="BDD7EE")
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
BASE_FONT = Font(name="游ゴシック", size=11)
HEAD_FONT = Font(name="游ゴシック", size=11, bold=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

INVALID_SHEET = re.compile(r"[\[\]\:\*\?\/\\]")


def sheet_name(base: str, used: set) -> str:
    name = INVALID_SHEET.sub("", base)[:31] or "sheet"
    candidate, i = name, 2
    while candidate in used:
        suffix = f"_{i}"
        candidate = name[: 31 - len(suffix)] + suffix
        i += 1
    used.add(candidate)
    return candidate


def _write_sheet(ws, fields, rows: List[dict], decorate: bool) -> None:
    ncol = len(fields)

    # 1〜3行目：見出し
    for i, fd in enumerate(fields, start=1):
        ws.cell(1, i, fd.group1 or None)
        ws.cell(2, i, fd.group2 or None)
        ws.cell(3, i, fd.column)

    if decorate:
        for r, attr in ((1, "group1"), (2, "group2")):
            i = 1
            while i <= ncol:
                val = getattr(fields[i - 1], attr)
                if not val:
                    i += 1
                    continue
                j = i
                while j < ncol and getattr(fields[j], attr) == val:
                    j += 1
                if j > i:
                    ws.merge_cells(start_row=r, start_column=i, end_row=r, end_column=j)
                i = j + 1

    for r in (1, 2, 3):
        for c in range(1, ncol + 1):
            cell = ws.cell(r, c)
            cell.font = HEAD_FONT
            cell.alignment = CENTER
            if decorate:
                cell.fill = GROUP_FILL if r < 3 else HEADER_FILL
                cell.border = BORDER

    # 4行目以降：データ
    for ri, row in enumerate(rows, start=4):
        for ci, fd in enumerate(fields, start=1):
            cell = ws.cell(ri, ci, row.get(fd.column, ""))
            cell.font = BASE_FONT
            if decorate:
                cell.border = BORDER
            if fd.transform == "int":
                cell.alignment = Alignment(horizontal="right")

    # 列幅
    for ci, fd in enumerate(fields, start=1):
        width = max(
            [len(str(fd.column)) * 2 + 2]
            + [min(len(str(r.get(fd.column, ""))) * 1.6 + 2, 50) for r in rows[:300]]
        )
        ws.column_dimensions[get_column_letter(ci)].width = max(9.5, min(width, 50))

    if decorate:
        ws.freeze_panes = "B4"
        if rows:
            ws.auto_filter.ref = f"A3:{get_column_letter(ncol)}{3 + len(rows)}"


def write_workbook(path: str, services: Dict[str, object], data: Dict[str, List[dict]],
                   summary_sheet: bool, per_city_sheet: bool, decorate: bool,
                   meta: Dict[str, object] | None = None) -> str:
    """data: {サービス名: [行dict, ...]} を1ブックに書き出す。"""
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()

    for svc_name, sd in services.items():
        rows = data.get(svc_name, [])
        if summary_sheet:
            ws = wb.create_sheet(sheet_name(f"{sd.short_name}_全件", used))
            _write_sheet(ws, sd.fields, rows, decorate)
        if per_city_sheet:
            by_city: Dict[str, List[dict]] = {}
            for r in rows:
                by_city.setdefault(r.get("市区町村", "不明"), []).append(r)
            for city, crows in by_city.items():
                ws = wb.create_sheet(sheet_name(f"{sd.short_name}_{city}", used))
                _write_sheet(ws, sd.fields, crows, decorate)

    # 実行サマリ
    ws = wb.create_sheet(sheet_name("実行サマリ", used))
    ws.append(["項目", "内容"])
    for k, v in (meta or {}).items():
        ws.append([k, str(v)])
    ws.append(["", ""])
    ws.append(["サービス", "取得件数"])
    for svc_name in services:
        ws.append([svc_name, len(data.get(svc_name, []))])
    for row in ws.iter_rows():
        for cell in row:
            cell.font = BASE_FONT
    ws.column_dimensions["A"].width = 28
    ws.column_dimensions["B"].width = 70

    if not wb.sheetnames:
        wb.create_sheet("結果なし")

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    try:
        wb.save(path)
    except PermissionError:
        alt = re.sub(r"\.xlsx$", "", path) + "_1.xlsx"
        log.error("出力先を開いているため保存できません。%s に保存します", alt)
        wb.save(alt)
        return alt
    return path
