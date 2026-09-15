"""画面版（gui.py）を実際に起動し、実行から出力までを通す検証。

tkinter が無い環境では自動でスキップする。
（Linuxでは python3-tk が必要。Windows の Python には最初から入っている）
"""
import os
import shutil
import sys
import tempfile
import time
import unittest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
sys.path.insert(0, os.path.join(BASE, "tests"))


def gui_available():
    try:
        import tkinter  # noqa: F401
        import customtkinter  # noqa: F401
    except ImportError:
        return False
    return bool(os.environ.get("DISPLAY"))


@unittest.skipUnless(gui_available(), "tkinter/customtkinter/画面が無いためスキップ")
class TestGuiRun(unittest.TestCase):
    def setUp(self):
        from mocksite import MockSite

        self.work = tempfile.mkdtemp()
        shutil.copytree(os.path.join(BASE, "config"), os.path.join(self.work, "config"))
        shutil.copy2(os.path.join(BASE, "settings.ini"), self.work)
        cache = os.path.join(self.work, "config", "cities_cache")
        os.makedirs(cache, exist_ok=True)
        with open(os.path.join(cache, "28.json"), "w", encoding="utf-8") as f:
            f.write('["西宮市", "芦屋市", "神戸市中央区"]')

        self.site = MockSite("real")
        self.site.__enter__()
        os.environ["KAIGO_BASE_URL"] = self.site.base_url

        import gui
        gui.messagebox.askokcancel = lambda *a, **k: True   # 確認ダイアログは自動でOK
        self.gui = gui
        gui.BASE_DIR = self.work
        gui.CACHE_DIR = cache
        self.app = gui.App()

    def tearDown(self):
        try:
            self.app._on_close()
        except Exception:
            pass
        self.site.__exit__(None, None, None)
        os.environ.pop("KAIGO_BASE_URL", None)
        shutil.rmtree(self.work, ignore_errors=True)

    def pump(self, seconds):
        end = time.time() + seconds
        while time.time() < end:
            self.app.update()
            time.sleep(0.05)

    def wait_until(self, cond, timeout):
        end = time.time() + timeout
        while time.time() < end:
            self.app.update()
            if cond():
                return True
            time.sleep(0.1)
        return False

    def test_screen_is_built(self):
        self.assertEqual(len(self.app.city_vars), 3)
        self.assertEqual(set(self.app.svc_vars), {"居宅介護支援", "訪問看護"})
        self.assertEqual(self.app.run_btn.cget("state"), "normal")
        self.assertEqual(self.app.stop_btn.cget("state"), "disabled")

    def test_select_all_and_clear(self):
        self.app._set_all(True)
        self.assertEqual(len(self.app.selected_cities()), 3)
        self.app._set_all(False)
        self.assertEqual(self.app.selected_cities(), [])

    def test_select_wards_only(self):
        self.app._set_all(False)
        self.app._select_wards()
        self.assertEqual(self.app.selected_cities(), ["神戸市中央区"])

    def test_filter_then_select_visible(self):
        """絞り込んだうえで「表示中だけ選ぶ」が効くこと。

        既定の「すべて選択／解除」は絞り込みに関係なく全件に効く。
        画面外の選択が残ったまま実行される事故を防ぐため。
        """
        self.app._set_all(False)
        self.app.filter_entry.delete(0, "end")
        self.app.filter_entry.insert(0, "神戸市")
        self.app._render_cities()
        self.app._set_all(True, visible_only=True)
        self.assertEqual(self.app.selected_cities(), ["神戸市中央区"])

    def test_confirm_dialog_blocks_run(self):
        """確認でキャンセルしたら実行しないこと。"""
        self.gui.messagebox.askokcancel = lambda *a, **k: False
        self.app._set_all(False)
        self.app.city_vars["芦屋市"].set(True)
        self.app._start()
        self.assertEqual(self.app.run_btn.cget("state"), "normal", "キャンセルしたのに走り出した")
        self.gui.messagebox.askokcancel = lambda *a, **k: True

    def test_select_all_ignores_filter(self):
        """絞り込み中でも「すべて解除」は全件に効くこと（画面外の選択が残らない）。"""
        self.app._set_all(True)
        self.app.filter_entry.delete(0, "end")
        self.app.filter_entry.insert(0, "芦屋")
        self.app._render_cities()
        self.app._set_all(False)
        self.assertEqual(self.app.selected_cities(), [], "画面外の選択が残っている")

    def test_warning_shown_next_to_button(self):
        self.app._set_all(False)
        self.app._start()
        self.assertIn("市区町村", self.app.warn_label.cget("text"))

    def test_trial_checkbox(self):
        """お試しのチェックで件数上限が渡ること。"""
        self.app.trial_var.set(False)
        self.assertEqual(self.app.trial_limit(), 0)
        self.app.trial_var.set(True)
        self.app.trial_n.configure(state="normal")
        self.app.trial_n.delete(0, "end")
        self.app.trial_n.insert(0, "3")
        self.assertEqual(self.app.trial_limit(), 3)
        self.app.trial_n.delete(0, "end")       # 空欄や不正値でも落ちないこと
        self.assertEqual(self.app.trial_limit(), 5)

    def test_trial_run_makes_separate_file(self):
        from openpyxl import load_workbook

        self.app._set_all(False)
        self.app.city_vars["西宮市"].set(True)
        self.app.svc_vars["訪問看護"].set(False)
        self.app.trial_var.set(True)
        self.app.trial_n.configure(state="normal")
        self.app.trial_n.delete(0, "end")
        self.app.trial_n.insert(0, "2")
        self.app._start()
        self.assertTrue(
            self.wait_until(lambda: self.app.run_btn.cget("state") == "normal", 180),
            "お試し実行が終わらない",
        )
        self.assertIn("お試し", os.path.basename(self.app.last_output))
        wb = load_workbook(self.app.last_output)
        self.assertEqual(wb["居宅_全件"].max_row, 3 + 2)
        self.assertIn("お試し", self.app.status.cget("text"))

    def test_run_produces_excel(self):
        from openpyxl import load_workbook

        self.app._set_all(False)
        self.app.city_vars["芦屋市"].set(True)
        self.app.svc_vars["訪問看護"].set(False)
        self.app.trial_var.set(False)
        self.app._start()
        self.assertEqual(self.app.run_btn.cget("state"), "disabled")
        self.assertTrue(
            self.wait_until(lambda: self.app.run_btn.cget("state") == "normal", 180),
            "実行が終わらない",
        )
        self.assertTrue(self.app.last_output, "出力ファイルが記録されていない")
        self.assertTrue(os.path.exists(self.app.last_output))
        self.assertIn("兵庫県", os.path.basename(self.app.last_output))
        wb = load_workbook(self.app.last_output)
        self.assertIn("居宅_全件", wb.sheetnames)
        self.assertEqual(wb["居宅_全件"].max_row, 3 + 3)
        self.assertEqual(wb["居宅_全件"].cell(4, 1).value, "芦屋市")
        self.assertGreater(self.app.bar.get(), 0.9)
        self.assertIn("完了しました", self.app.status.cget("text"))

    def test_stop_button(self):
        self.app._set_all(True)
        self.app._start()
        self.pump(6)                      # 走り出してから止める
        self.app._stop()
        self.assertTrue(
            self.wait_until(lambda: self.app.run_btn.cget("state") == "normal", 120),
            "停止できない",
        )
        self.assertIn("停止", self.app.status.cget("text"))


@unittest.skipUnless(gui_available(), "tkinter/customtkinter/画面が無いためスキップ")
class TestFirstLaunch(unittest.TestCase):
    """市区町村のキャッシュが無い状態（初めて使う人が必ず通る経路）で起動できること。

    ここを試していなかったため、「実行しても画面が出ない」不具合を出した。
    キャッシュがあるかないかで通る道が変わるので、両方を試すこと。
    """

    def setUp(self):
        self.work = tempfile.mkdtemp()
        shutil.copytree(os.path.join(BASE, "config"), os.path.join(self.work, "config"))
        shutil.copy2(os.path.join(BASE, "settings.ini"), self.work)
        # cities_cache は作らない（＝初回起動）
        self.cache = os.path.join(self.work, "config", "cities_cache")
        import gui
        self.gui = gui
        gui.BASE_DIR = self.work
        gui.CACHE_DIR = self.cache
        # 市区町村の取得でサイトへ行かないようにする（起動できるかだけを見る）
        self._orig = gui.App._fetch_cities
        gui.App._fetch_cities = lambda self, pref, code: None

    def tearDown(self):
        self.gui.App._fetch_cities = self._orig
        shutil.rmtree(self.work, ignore_errors=True)

    def test_starts_without_city_cache(self):
        app = self.gui.App()
        try:
            app.update()
            self.assertEqual(app.all_cities, [])
            self.assertIn("読み込", app._empty_label.cget("text"))
            self.assertEqual(app.run_btn.cget("state"), "disabled", "読み込み中は実行できない")
        finally:
            app._on_close()

    def test_recovers_when_city_fetch_fails(self):
        """市区町村の取得に失敗しても、画面が固まらず操作に戻れること。"""
        app = self.gui.App()
        try:
            app.events.put(("log", "市区町村の取得に失敗しました: 通信エラー"))
            app.events.put(("cities", []))
            end = time.time() + 10
            while time.time() < end and app.run_btn.cget("state") != "normal":
                app.update()
                time.sleep(0.05)
            self.assertEqual(app.run_btn.cget("state"), "normal", "操作に戻れない")
            # 「読み込み中」のままにせず、失敗したことと次の一手を示すこと
            self.assertIn("取得できませんでした", app._empty_label.cget("text"))
            self.assertIn("取り直す", app._empty_label.cget("text"))
        finally:
            app._on_close()


@unittest.skipUnless(gui_available(), "tkinter/customtkinter/画面が無いためスキップ")
class TestCityFetchFallback(unittest.TestCase):
    """非表示で取れなかったら、ブラウザを表示して取り直すこと。

    実サイトでは、ブラウザ非表示のときに市区町村が1件も取れない事象が起きた。
    利用者に settings.ini を触らせず、ツール側で復旧する。
    """

    def setUp(self):
        self.work = tempfile.mkdtemp()
        shutil.copytree(os.path.join(BASE, "config"), os.path.join(self.work, "config"))
        with open(os.path.join(BASE, "settings.ini"), encoding="utf-8-sig") as f:
            ini = f.read().replace("display = True", "display = False")
        with open(os.path.join(self.work, "settings.ini"), "w", encoding="utf-8") as f:
            f.write(ini)
        self.cache = os.path.join(self.work, "config", "cities_cache")
        import gui
        self.gui = gui
        gui.BASE_DIR = self.work
        gui.CACHE_DIR = self.cache
        self.calls = []

        class FakeNav:
            def __init__(inner, display):
                inner.display = display

            def list_cities(inner, label):
                self.calls.append(inner.display)
                return ["西宮市", "芦屋市"] if inner.display else []

            def close(inner):
                pass

        self._orig = gui.runner.make_navigator
        gui.runner.make_navigator = lambda s, code: FakeNav(s.display)
        # 起動時の自動取得は止め、この検証だけを走らせる
        self._orig_load = gui.App._load_cities
        gui.App._load_cities = lambda self, initial=False, force=False: None

    def tearDown(self):
        self.gui.runner.make_navigator = self._orig
        self.gui.App._load_cities = self._orig_load
        shutil.rmtree(self.work, ignore_errors=True)

    def test_retries_with_visible_browser(self):
        app = self.gui.App()
        try:
            app._fetch_cities("兵庫県", "28")
            self.assertEqual(self.calls, [False, True], "表示ありで取り直していない")
            kind, payload = app.events.get_nowait(), None
            while kind[0] != "cities":
                kind = app.events.get_nowait()
            self.assertEqual(kind[1], ["西宮市", "芦屋市"])
            # 次回からは最初から表示ありで動くよう、設定に書き戻すこと
            with open(os.path.join(self.work, "settings.ini"), encoding="utf-8-sig") as f:
                self.assertIn("display = True", f.read())
        finally:
            app._on_close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
