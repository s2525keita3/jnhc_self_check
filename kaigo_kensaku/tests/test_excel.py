"""抽出 → Excel出力までの結合テスト（ネットワーク不要）。"""
import os
import sys
import tempfile
import unittest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)

from openpyxl import load_workbook  # noqa: E402

from src.config import load_services  # noqa: E402
from src.excelout import write_workbook  # noqa: E402
from src.extract import build_row, harvest_html  # noqa: E402


def rows_from(fixture, sd, city, service):
    with open(os.path.join(BASE, "tests", "fixtures", fixture), encoding="utf-8") as f:
        h = harvest_html(f.read())
    ctx = {
        "city": city,
        "service": service,
        "pref": "兵庫県",
        "name": f"テスト{service}事業所",
        "jigyosyo_cd": "2800000000-00",
        "url": "https://example.invalid/x",
    }
    return [build_row(sd.fields, h, ctx, normalize=True)]


class TestWorkbook(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.services = load_services(BASE)
        data = {
            "居宅介護支援": rows_from("sample_detail.html", cls.services["居宅介護支援"], "西宮市", "居宅介護支援")
            + rows_from("sample_detail.html", cls.services["居宅介護支援"], "芦屋市", "居宅介護支援"),
            "訪問看護": rows_from("sample_houkan.html", cls.services["訪問看護"], "神戸市中央区", "訪問看護"),
        }
        cls.tmp = tempfile.mkdtemp()
        cls.path = write_workbook(
            os.path.join(cls.tmp, "out.xlsx"),
            cls.services,
            data,
            summary_sheet=True,
            per_city_sheet=True,
            decorate=True,
            meta={"実行日時": "2026/01/30 15:26", "都道府県": "兵庫県"},
        )
        cls.wb = load_workbook(cls.path)

    def test_sheets_c_plan(self):
        names = self.wb.sheetnames
        # C案：サービス別の統合シート ＋ 市区町村別シート ＋ 実行サマリ
        self.assertIn("居宅_全件", names)
        self.assertIn("訪看_全件", names)
        self.assertIn("居宅_西宮市", names)
        self.assertIn("居宅_芦屋市", names)
        self.assertIn("訪看_神戸市中央区", names)
        self.assertIn("実行サマリ", names)

    def test_header_is_three_rows(self):
        ws = self.wb["居宅_全件"]
        self.assertEqual(ws.cell(3, 1).value, "市区町村")
        self.assertEqual(ws.cell(3, 3).value, "事業所名")
        self.assertEqual(ws.cell(1, 8).value, "営業時間")
        self.assertEqual(ws.cell(2, 15).value, "要介護度別利用者数")
        self.assertEqual(ws.cell(3, 23).value, "（Ⅰ）")

    def test_summary_contains_all_cities(self):
        ws = self.wb["居宅_全件"]
        self.assertEqual(ws.max_row, 5)  # ヘッダー3行 + 2件
        self.assertEqual([ws.cell(r, 1).value for r in (4, 5)], ["西宮市", "芦屋市"])

    def test_per_city_sheet_has_only_that_city(self):
        ws = self.wb["居宅_芦屋市"]
        self.assertEqual(ws.max_row, 4)
        self.assertEqual(ws.cell(4, 1).value, "芦屋市")

    def test_houkan_columns(self):
        ws = self.wb["訪看_全件"]
        heads = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        self.assertIn("看護師_常勤", heads)
        self.assertIn("緊急時訪問看護加算", heads)
        row = {heads[c - 1]: ws.cell(4, c).value for c in range(1, ws.max_column + 1)}
        self.assertEqual(row["看護師_常勤"], 5)
        self.assertEqual(row["看護師_非常勤"], 3)
        self.assertEqual(row["緊急時訪問看護加算"], "あり")
        self.assertEqual(row["事業開始年月日"], "2021/04/01")
        self.assertEqual(row["所在地"], "兵庫県神戸市中央区加納町6-5-1")
        self.assertEqual(row["土曜"], None)

    def test_decoration(self):
        ws = self.wb["居宅_全件"]
        self.assertEqual(ws.freeze_panes, "B4")
        self.assertTrue(ws.auto_filter.ref.startswith("A3:"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
