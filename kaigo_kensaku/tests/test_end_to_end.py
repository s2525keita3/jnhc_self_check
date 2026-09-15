"""main.py をそのまま実行し、Excelが出来るところまで確認する結合テスト。"""
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE)
sys.path.insert(0, os.path.join(BASE, "tests"))

from mocksite import MockSite  # noqa: E402
from test_navigator_mock import browser_available  # noqa: E402

SETTINGS = """[dev]
debug = True
display = False
dump_html = False

[search]
pref = 兵庫県
cities = {cities}
search_type = partial
services = {services}

[output]
file = 出力.xlsx
summary_sheet = True
per_city_sheet = True
decorate = True
normalize = True

[run]
wait = 0
retry = 2
resume = {resume}
"""


@unittest.skipUnless(browser_available(), "Chrome/chromedriver が無いためスキップ")
class TestEndToEnd(unittest.TestCase):
    def setUp(self):
        self.work = tempfile.mkdtemp()
        for item in ("main.py", "src", "config"):
            src = os.path.join(BASE, item)
            dst = os.path.join(self.work, item)
            if os.path.isdir(src):
                shutil.copytree(src, dst, ignore=shutil.ignore_patterns("__pycache__"))
            else:
                shutil.copy2(src, dst)

    def tearDown(self):
        shutil.rmtree(self.work, ignore_errors=True)

    def _write_settings(self, cities, services, resume="True"):
        with open(os.path.join(self.work, "settings.ini"), "w", encoding="utf-8") as f:
            f.write(SETTINGS.format(cities=cities, services=services, resume=resume))

    def _run(self, site, extra=()):
        env = dict(os.environ, KAIGO_BASE_URL=site.base_url, PYTHONIOENCODING="utf-8")
        return subprocess.run(
            [sys.executable, "main.py", *extra],
            cwd=self.work, env=env, capture_output=True, text=True, timeout=600,
            stdin=subprocess.DEVNULL,
        )

    def test_full_run_two_services_multi_city(self):
        from openpyxl import load_workbook

        self._write_settings("西宮市, 芦屋市", "居宅介護支援, 訪問看護")
        with MockSite("checkbox") as site:
            r = self._run(site)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("完了：", r.stdout)

        out = os.path.join(self.work, "出力.xlsx")
        self.assertTrue(os.path.exists(out), r.stdout)
        wb = load_workbook(out)
        # C案：統合シート＋市区町村別シート＋実行サマリ
        for name in ("居宅_全件", "居宅_西宮市", "居宅_芦屋市",
                     "訪看_全件", "訪看_西宮市", "訪看_芦屋市", "実行サマリ"):
            self.assertIn(name, wb.sheetnames)
        ws = wb["居宅_全件"]
        self.assertEqual(ws.max_row, 3 + 12 + 3)      # 西宮市12件 + 芦屋市3件
        self.assertEqual(ws.cell(3, 1).value, "市区町村")
        self.assertEqual({ws.cell(r, 1).value for r in range(4, ws.max_row + 1)},
                         {"西宮市", "芦屋市"})
        self.assertEqual(wb["居宅_芦屋市"].max_row, 4 + 2)
        self.assertEqual(wb["訪看_全件"].max_row, 3 + 3 + 3)
        # 値が実際に入っていること
        heads = [ws.cell(3, c).value for c in range(1, ws.max_column + 1)]
        row = {heads[c - 1]: ws.cell(4, c).value for c in range(1, ws.max_column + 1)}
        self.assertTrue(row["事業所名"])
        self.assertTrue(str(row["所在地"]).startswith("兵庫県"))
        self.assertEqual(row["（Ⅱ）"], "あり")
        # 空欄だらけの列が無いこと
        summary = wb["実行サマリ"]
        texts = [str(c.value) for r in summary.iter_rows() for c in r]
        self.assertNotIn("要確認列（居宅介護支援）", texts)

    def test_real_flow_two_services(self):
        """実サイトの導線（サービス→所在地）でも最後まで通ること。"""
        from openpyxl import load_workbook

        self._write_settings("西宮市, 芦屋市", "居宅介護支援, 訪問看護")
        with MockSite("real") as site:
            r = self._run(site)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        wb = load_workbook(os.path.join(self.work, "出力.xlsx"))
        self.assertEqual(wb["居宅_全件"].max_row, 3 + 12 + 3)
        self.assertEqual(wb["訪看_全件"].max_row, 3 + 3 + 3)
        self.assertIn("訪看_芦屋市", wb.sheetnames)

    def test_real_flow_wildcard_wards(self):
        """実サイトの導線で、神戸市* の展開（市区町村一覧の取得）が効くこと。"""
        from openpyxl import load_workbook

        self._write_settings("神戸市*", "居宅介護支援")
        with MockSite("real") as site:
            r = self._run(site)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        wb = load_workbook(os.path.join(self.work, "出力.xlsx"))
        for ward in ("神戸市東灘区", "神戸市灘区", "神戸市中央区"):
            self.assertIn(f"居宅_{ward}", wb.sheetnames)

    def test_wildcard_expands_wards(self):
        from openpyxl import load_workbook

        self._write_settings("神戸市*", "居宅介護支援")
        with MockSite("checkbox") as site:
            r = self._run(site)
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        wb = load_workbook(os.path.join(self.work, "出力.xlsx"))
        for ward in ("神戸市東灘区", "神戸市灘区", "神戸市中央区"):
            self.assertIn(f"居宅_{ward}", wb.sheetnames)
        self.assertEqual(wb["居宅_全件"].max_row, 3 + 9)   # 3区 × 3件

    def test_resume_skips_completed(self):
        self._write_settings("西宮市", "居宅介護支援")
        with MockSite("checkbox") as site:
            first = self._run(site)
            self.assertEqual(first.returncode, 0, first.stdout + first.stderr)
            second = self._run(site)
        self.assertIn("取得済みのためスキップ", second.stdout)
        self.assertIn("完了：12件", second.stdout)

    def test_list_cities_option(self):
        self._write_settings("西宮市", "居宅介護支援")
        with MockSite("checkbox") as site:
            r = self._run(site, ["--list-cities"])
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("神戸市中央区", r.stdout)

    def test_diagnose_option(self):
        self._write_settings("西宮市", "居宅介護支援")
        with MockSite("link") as site:
            r = self._run(site, ["--diagnose"])
        self.assertEqual(r.returncode, 0, r.stdout + r.stderr)
        self.assertIn("リンク", r.stdout)
        self.assertTrue(os.path.exists(os.path.join(self.work, "logs", "診断結果.txt")))


if __name__ == "__main__":
    unittest.main(verbosity=2)
